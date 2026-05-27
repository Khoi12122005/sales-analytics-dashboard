from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from src.data_loader import create_mysql_engine


HEADER_FILL = PatternFill("solid", fgColor="0F766E")
HEADER_FONT = Font(color="FFFFFF", bold=True)
BODY_ALIGNMENT = Alignment(vertical="center")


def _quality_log_to_df(quality_log: dict[str, Any]) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for key, value in quality_log.items():
        if isinstance(value, dict):
            for sub_key, sub_value in value.items():
                rows.append({"metric": f"{key}.{sub_key}", "value": sub_value})
        elif isinstance(value, list):
            if value:
                rows.append({"metric": key, "value": " | ".join(str(item) for item in value)})
            else:
                rows.append({"metric": key, "value": ""})
        else:
            rows.append({"metric": key, "value": value})
    return pd.DataFrame(rows)


def _build_kpi_summary_df(kpis: dict[str, Any]) -> pd.DataFrame:
    growth = kpis.get("monthly_growth_rate")
    growth_value = None if pd.isna(growth) else float(growth)

    return pd.DataFrame(
        [
            {"kpi": "Total Revenue", "value": float(kpis.get("total_revenue", 0.0)), "unit": "currency"},
            {"kpi": "Total Orders", "value": int(kpis.get("total_orders", 0)), "unit": "count"},
            {
                "kpi": "Average Order Value",
                "value": float(kpis.get("average_order_value", 0.0)),
                "unit": "currency",
            },
            {
                "kpi": "Total Quantity Sold",
                "value": int(kpis.get("total_quantity_sold", 0)),
                "unit": "count",
            },
            {"kpi": "Monthly Growth Rate", "value": growth_value, "unit": "percentage"},
        ]
    )


def _add_excel_table(ws, table_name: str) -> None:
    if ws.max_row < 2 or ws.max_column < 1:
        return

    ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    table = Table(displayName=table_name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def _format_sheet(ws) -> None:
    if ws.max_row >= 1:
        for cell in ws[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.alignment = BODY_ALIGNMENT

    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        max_len = 0
        for cell in ws[letter]:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(value))
        ws.column_dimensions[letter].width = min(max(12, max_len + 2), 45)


def _apply_number_formats(ws, sheet_name: str) -> None:
    if sheet_name == "KPI_Summary":
        value_col = None
        unit_col = None
        for col_idx in range(1, ws.max_column + 1):
            header = str(ws.cell(row=1, column=col_idx).value or "").lower()
            if header == "value":
                value_col = col_idx
            elif header == "unit":
                unit_col = col_idx

        if value_col and unit_col:
            for row_idx in range(2, ws.max_row + 1):
                value_cell = ws.cell(row=row_idx, column=value_col)
                unit = str(ws.cell(row=row_idx, column=unit_col).value or "").lower()
                if not isinstance(value_cell.value, (int, float)):
                    continue
                if unit == "currency":
                    value_cell.number_format = "$#,##0.00"
                elif unit == "percentage":
                    value_cell.number_format = "0.00%"
                elif unit == "count":
                    value_cell.number_format = "#,##0"
        return

    headers = [ws.cell(row=1, column=col_idx).value for col_idx in range(1, ws.max_column + 1)]
    lower_headers = [str(h).lower() if h is not None else "" for h in headers]

    for row_idx in range(2, ws.max_row + 1):
        for col_idx, header in enumerate(lower_headers, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if cell.value is None or not isinstance(cell.value, (int, float)):
                continue
            if any(token in header for token in ["revenue", "price", "value"]):
                cell.number_format = "$#,##0.00"
            elif any(token in header for token in ["growth", "rate", "share", "percentage"]):
                cell.number_format = "0.00%"
            elif any(token in header for token in ["quantity", "orders", "count"]):
                cell.number_format = "#,##0"


def export_summary_report(
    clean_df: pd.DataFrame,
    kpis: dict[str, Any],
    charts_data: dict[str, Any] | None,
    quality_log: dict[str, Any],
    output_dir: str = "reports",
) -> str:
    """Export a professional Excel report and return the generated file path."""
    del charts_data  # Reserved for future chart-image embedding support.

    output_path = Path(output_dir)
    output_path.mkdir(parents=True, exist_ok=True)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    report_file = output_path / f"sales_summary_report_{timestamp}.xlsx"

    executive_summary_df = pd.DataFrame(
        {
            "summary_item": [
                "Report Generated At",
                "Rows in Cleaned Dataset",
                "Total Revenue",
                "Total Orders",
                "Business Insights",
            ],
            "value": [
                datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                int(len(clean_df)),
                float(kpis.get("total_revenue", 0.0)),
                int(kpis.get("total_orders", 0)),
                " | ".join(kpis.get("insights", [])),
            ],
        }
    )

    kpi_summary_df = _build_kpi_summary_df(kpis)
    revenue_by_month_df = kpis.get("revenue_by_month", pd.DataFrame())
    revenue_by_category_df = kpis.get("revenue_by_category", pd.DataFrame())
    top_products_df = kpis.get("top_products", pd.DataFrame())
    top_customers_df = kpis.get("top_customers", pd.DataFrame())
    quality_log_df = _quality_log_to_df(quality_log)

    with pd.ExcelWriter(report_file, engine="openpyxl") as writer:
        executive_summary_df.to_excel(writer, sheet_name="Executive_Summary", index=False)
        kpi_summary_df.to_excel(writer, sheet_name="KPI_Summary", index=False)
        revenue_by_month_df.to_excel(writer, sheet_name="Revenue_By_Month", index=False)
        revenue_by_category_df.to_excel(writer, sheet_name="Revenue_By_Category", index=False)
        top_products_df.to_excel(writer, sheet_name="Top_Products", index=False)
        top_customers_df.to_excel(writer, sheet_name="Top_Customers", index=False)
        clean_df.to_excel(writer, sheet_name="Cleaned_Data", index=False)
        quality_log_df.to_excel(writer, sheet_name="Data_Quality_Log", index=False)

        workbook = writer.book
        table_counter = 1
        for sheet_name in workbook.sheetnames:
            ws = workbook[sheet_name]
            _format_sheet(ws)
            _apply_number_formats(ws, sheet_name)
            _add_excel_table(ws, f"Table{table_counter}")
            table_counter += 1

    return str(report_file)


def save_to_mysql(
    clean_df: pd.DataFrame,
    kpi_summary: dict[str, Any],
    db_config: dict[str, str],
) -> tuple[bool, str]:
    """Persist cleaned sales data and KPI snapshot to MySQL."""
    try:
        engine = create_mysql_engine(db_config)

        clean_df_to_save = clean_df.copy()
        clean_df_to_save["snapshot_at"] = datetime.now()
        clean_df_to_save.to_sql("sales_cleaned", con=engine, if_exists="append", index=False)

        kpi_df = pd.DataFrame(
            [
                {
                    "snapshot_at": datetime.now(),
                    "total_revenue": float(kpi_summary.get("total_revenue", 0.0)),
                    "total_orders": int(kpi_summary.get("total_orders", 0)),
                    "average_order_value": float(kpi_summary.get("average_order_value", 0.0)),
                    "total_quantity_sold": int(kpi_summary.get("total_quantity_sold", 0)),
                    "monthly_growth_rate": None
                    if pd.isna(kpi_summary.get("monthly_growth_rate"))
                    else float(kpi_summary.get("monthly_growth_rate")),
                }
            ]
        )
        kpi_df.to_sql("kpi_summary", con=engine, if_exists="append", index=False)

        return True, "Data persisted to MySQL successfully (tables: sales_cleaned, kpi_summary)."
    except Exception as exc:
        error_text = str(exc).lower()
        if "access denied" in error_text or "1045" in error_text:
            return False, "MySQL save skipped: access denied. Please verify DB_USER/DB_PASSWORD in .env."
        if "can't connect" in error_text or "connection refused" in error_text:
            return False, "MySQL save skipped: cannot connect to MySQL server. Please start MySQL and retry."
        if "unknown database" in error_text or "1049" in error_text:
            return False, "MySQL save skipped: database not found and could not be created automatically."
        return False, f"MySQL save skipped: {exc.__class__.__name__}. Check database configuration in .env."
