from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from src.kpi_calculator import calculate_kpis
from src.report_exporter import export_summary_report


def test_export_summary_report_creates_professional_workbook(tmp_path: Path):
    clean_df = pd.DataFrame(
        {
            "order_id": ["O1", "O2"],
            "order_date": pd.to_datetime(["2026-03-01", "2026-03-02"]),
            "customer_name": ["Alice", "Bob"],
            "product_name": ["Laptop", "Mouse"],
            "category": ["Electronics", "Accessories"],
            "quantity": [1, 2],
            "unit_price": [1000.0, 25.0],
            "revenue": [1000.0, 50.0],
        }
    )

    kpis = calculate_kpis(clean_df)
    quality_log = {
        "rows_initial": 2,
        "rows_final": 2,
        "invalid_date_rows_removed": 0,
        "duplicate_rows_removed": 0,
        "generated_order_ids": 0,
        "revenue_recalculated_rows": 0,
        "missing_columns_added": [],
        "missing_values_filled": {"quantity": 0},
        "notes": ["Test note"],
    }

    output_file = export_summary_report(
        clean_df=clean_df,
        kpis=kpis,
        charts_data={},
        quality_log=quality_log,
        output_dir=str(tmp_path),
    )

    assert Path(output_file).exists()

    workbook = load_workbook(output_file)
    expected_sheets = {
        "Executive_Summary",
        "KPI_Summary",
        "Revenue_By_Month",
        "Revenue_By_Category",
        "Top_Products",
        "Top_Customers",
        "Cleaned_Data",
        "Data_Quality_Log",
    }
    assert expected_sheets.issubset(set(workbook.sheetnames))
